from pathlib import Path
import json, sys
import openpyxl

paths = [Path(p) for p in sys.argv[1:]]
for path in paths:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("\nFILE", path.name)
    for ws in wb.worksheets:
        ws.reset_dimensions()
        print("SHEET", ws.title, "rows", ws.max_row, "cols", ws.max_column)
        rows = ws.iter_rows(values_only=True)
        for i, row in enumerate(rows, 1):
            vals = [None if v is None else str(v)[:160] for v in row]
            if any(v not in (None, "") for v in vals):
                print("ROW", i, json.dumps(vals, ensure_ascii=False))
            if i >= 8:
                break
